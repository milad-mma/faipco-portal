"""
Endpoint های منابع انسانی برای پیام تبریک تولد:
/hr/birthday-templates      (GET/POST/DELETE) مجموعه متن‌های تبریک تولد
/hr/birthday-send-time      (GET/PUT)         ساعت ارسال روزانه
/hr/birthday-enabled        (GET/PUT)         فعال/غیرفعال کلی قابلیت
/hr/birthday-send-now       (POST)            ارسال فوری پیام‌های امروز
/hr/birthdays/export        (GET)             خروجی Excel متولدین به تفکیک ماه (سایت‌محور)
همه با مجوز hr.birthday_messages؛ Admin و hr-manager هر دو روی همان داده مشترک کار می‌کنند.
چون این تنظیمات و ارسال فوری روی پرسنل همه‌ی سایت‌ها اثر دارند، تغییر آن‌ها (POST/PUT/DELETE)
فقط با انتصاب سراسری این مجوز (یا superuser) مجاز است؛ خواندن با انتصاب سایتی هم ممکن است.
"""
import asyncio
import io
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.deps import require_permission
from app.core.site_access import get_sites_with_permission
from app.db.session import get_db
from app.models.employee import Employee
from app.models.site import Site
from app.models.user import User
from app.schemas.birthday_greetings import (
    BirthdayEnabledIn,
    BirthdayEnabledOut,
    BirthdaySendTimeIn,
    BirthdaySendTimeOut,
    BirthdayTemplateIn,
    BirthdayTemplateOut,
)
from app.services.birthday_greetings_service import BirthdayGreetingsService

router = APIRouter()


async def require_org_wide_birthday_manager(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission("hr.birthday_messages")),
) -> User:
    """
    Dependency: کاربر باید hr.birthday_messages را به‌صورت سراسری (یا superuser) داشته باشد؛
    انتصاب سایتی کافی نیست چون متن‌ها، ساعت ارسال و ارسال فوری برای همه‌ی سایت‌ها مشترک‌اند. در غیر این صورت 403.
    """
    if await get_sites_with_permission(db, user, "hr.birthday_messages") is not None:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="تنظیمات پیام تبریک تولد برای همه‌ی سایت‌ها مشترک است و فقط با مجوز برای همه‌ی سایت‌ها قابل تغییر است",
        )
    return user


@router.get("/birthday-templates", response_model=list[BirthdayTemplateOut])
async def list_birthday_templates(
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_permission("hr.birthday_messages")),
):
    """فهرست همه متن‌های تبریک تولد را برمی‌گرداند. مجوز: hr.birthday_messages."""
    return await BirthdayGreetingsService(db).list_templates()


@router.post("/birthday-templates", response_model=BirthdayTemplateOut, status_code=status.HTTP_201_CREATED)
async def add_birthday_template(
    payload: BirthdayTemplateIn,
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_org_wide_birthday_manager),
):
    """یک متن تبریک جدید اضافه می‌کند و آن را برمی‌گرداند. مجوز: hr.birthday_messages؛ خطای 400 برای متن نامعتبر."""
    try:
        return await BirthdayGreetingsService(db).add_template(payload.text)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.delete("/birthday-templates/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_birthday_template(
    template_id: int,
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_org_wide_birthday_manager),
):
    """یک متن تبریک را حذف می‌کند. مجوز: hr.birthday_messages؛ خطای 404 اگر پیدا نشود."""
    deleted = await BirthdayGreetingsService(db).delete_template(template_id)
    if not deleted:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="پیام یافت نشد")


@router.get("/birthday-send-time", response_model=BirthdaySendTimeOut)
async def get_birthday_send_time(
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_permission("hr.birthday_messages")),
):
    """ساعت و دقیقه ارسال روزانه پیام تبریک را برمی‌گرداند. مجوز: hr.birthday_messages."""
    hour, minute = await BirthdayGreetingsService(db).get_send_time()
    return BirthdaySendTimeOut(hour=hour, minute=minute)


@router.put("/birthday-send-time", response_model=BirthdaySendTimeOut)
async def update_birthday_send_time(
    payload: BirthdaySendTimeIn,
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_org_wide_birthday_manager),
):
    """
    ساعت ارسال روزانه را ذخیره و زمان‌بند را فوراً به‌روز می‌کند.
    مجوز: hr.birthday_messages؛ خطای 400 برای ساعت/دقیقه نامعتبر.
    """
    from app.core.scheduler import reschedule_birthday_send_time

    try:
        hour, minute = await BirthdayGreetingsService(db).set_send_time(payload.hour, payload.minute)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    # بدون Restart سرور، همان لحظه روی Job در حال اجرا اعمال می‌شود
    reschedule_birthday_send_time(hour, minute)
    return BirthdaySendTimeOut(hour=hour, minute=minute)


@router.get("/birthday-enabled", response_model=BirthdayEnabledOut)
async def get_birthday_enabled(
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_permission("hr.birthday_messages")),
):
    """وضعیت فعال/غیرفعال بودن ارسال خودکار تبریک تولد را برمی‌گرداند. مجوز: hr.birthday_messages."""
    enabled = await BirthdayGreetingsService(db).get_enabled()
    return BirthdayEnabledOut(enabled=enabled)


@router.put("/birthday-enabled", response_model=BirthdayEnabledOut)
async def update_birthday_enabled(
    payload: BirthdayEnabledIn,
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_org_wide_birthday_manager),
):
    """ارسال خودکار تبریک تولد را فعال/غیرفعال می‌کند و وضعیت جدید را برمی‌گرداند. مجوز: hr.birthday_messages."""
    enabled = await BirthdayGreetingsService(db).set_enabled(payload.enabled)
    return BirthdayEnabledOut(enabled=enabled)


@router.post("/birthday-send-now")
async def send_birthday_greetings_now(
    db: AsyncSession = Depends(get_db),
    _user=Depends(require_org_wide_birthday_manager),
):
    """
    پیام‌های تبریک تولد امروز را فوراً و بدون صبر تا ساعت زمان‌بندی‌شده ارسال می‌کند (مثلاً برای تست).
    مجوز: hr.birthday_messages. خروجی: تعداد پیام‌های ارسال‌شده + پیام متنی برای نمایش.
    """
    sent_count = await BirthdayGreetingsService(db).send_todays_birthday_greetings()
    return {
        "sent_count": sent_count,
        "message": f"{sent_count} پیام تبریک تولد فرستاده شد." if sent_count else "هیچ پیامی فرستاده نشد (یا امروز کسی تولد ندارد، یا فهرست خالی است، یا ارسال خودکار غیرفعال است).",
    }


JALALI_MONTHS = ("فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند")


def _birthdays_xlsx(rows_by_month: dict[int, list[dict]], months: list[int], show_site: bool) -> bytes:
    """فایل Excel متولدین: یک برگه‌ی «خلاصه» (تعداد هر ماه) و یک برگه برای هر ماه (راست‌به‌چپ)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="185E95")

    def _style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.sheet_view.rightToLeft = True

    summary = wb.active
    summary.title = "خلاصه"
    summary.append(["ماه", "تعداد متولدین"])
    for m in months:
        summary.append([JALALI_MONTHS[m - 1], len(rows_by_month.get(m, []))])
    summary.append(["جمع", sum(len(rows_by_month.get(m, [])) for m in months)])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    summary.cell(row=summary.max_row, column=2).font = Font(bold=True)
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 16
    _style_header(summary)

    headers = ["ردیف", "نام و نام خانوادگی", "کد پرسنلی", "واحد", "تاریخ تولد", "روز"] + (["سایت"] if show_site else [])
    widths = [7, 28, 14, 26, 14, 7, 18]
    for m in months:
        ws = wb.create_sheet(JALALI_MONTHS[m - 1])
        ws.append(headers)
        for i, r in enumerate(rows_by_month.get(m, []), start=1):
            ws.append([i, r["name"], r["code"], r["department"], r["birth_date"], r["day"]] + ([r["site"]] if show_site else []))
        for idx, width in enumerate(widths[: len(headers)], start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
        _style_header(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/birthdays/export")
async def export_birthdays(
    month: int | None = Query(default=None, ge=1, le=12),
    site_id: int | None = Query(default=None, ge=1),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_permission("hr.birthday_messages")),
):
    """
    خروجی Excel متولدین به تفکیک ماه شمسی: برگه‌ی «خلاصه» و یک برگه برای هر ماه (یا فقط ماه month) با
    نام و نام خانوادگی، کد پرسنلی، واحد و تاریخ تولد؛ مرتب بر اساس روز. فقط پرسنل فعال سایت‌هایی که کاربر
    برایشان hr.birthday_messages دارد؛ با site_id فقط همان سایت (سایت غیرمجاز → 403). (این گزارش داخلی است و تنظیم «نمایش تولد در داشبورد» پرسنل را در نظر نمی‌گیرد).
    """
    sites = await get_sites_with_permission(db, user, "hr.birthday_messages")  # None = همه‌ی سایت‌ها
    stmt = (
        select(Employee, Site.name)
        .join(Site, Site.id == Employee.site_id)
        .options(selectinload(Employee.department))
        .where(
            Employee.is_active.is_(True),
            Employee.is_enabled.is_(True),
            Employee.birth_month.is_not(None),
            Employee.birth_day.is_not(None),
        )
    )
    if site_id is not None:
        if sites is not None and site_id not in sites:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="به این سایت دسترسی ندارید.")
        stmt = stmt.where(Employee.site_id == site_id)
    elif sites is not None:
        stmt = stmt.where(Employee.site_id.in_(sites))
    if month is not None:
        stmt = stmt.where(Employee.birth_month == month)
    result = await db.execute(stmt)

    rows_by_month: dict[int, list[dict]] = {}
    site_names: set[str] = set()
    for emp, site_name in result.all():
        site_names.add(site_name)
        birth = emp.birth_date_jalali or f"{emp.birth_month:02d}/{emp.birth_day:02d}"
        rows_by_month.setdefault(emp.birth_month, []).append(
            {
                "name": f"{emp.first_name} {emp.last_name}".strip(),
                "code": emp.personnel_code,
                "department": emp.department.name if emp.department else "—",
                "birth_date": birth,
                "day": emp.birth_day,
                "site": site_name,
            }
        )
    for rows in rows_by_month.values():
        rows.sort(key=lambda r: (r["day"], r["name"]))

    months = [month] if month is not None else list(range(1, 13))
    content = await asyncio.to_thread(_birthdays_xlsx, rows_by_month, months, len(site_names) > 1)
    suffix = (f"-site{site_id}" if site_id is not None else "") + (f"-{month:02d}" if month is not None else "")
    filename = f"birthdays{suffix}-{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
