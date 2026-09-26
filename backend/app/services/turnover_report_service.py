"""
سرویس «گزارش جذب و ترک کار».

آمار استخدام و ترک کار مستقیماً از جدول پرسنل دیتابیس منبع (کاراوب: Employee) خوانده
می‌شود، چون پرسنل قطع‌همکاری‌شده وارد پرتال نمی‌شوند. همه‌ی نام جدول/ستون‌ها از نگاشت
پرسنل سایت (EmployeeMapping) می‌آیند. برای حریم خصوصی فقط ستون‌های غیرشناسایی‌کننده
خوانده می‌شوند (تاریخ‌ها، وضعیت، علت، واحد، جنسیت، تاریخ تولد، سمت، مدرک) — نه نام، نه
کد پرسنلی، نه کد ملی — و خروجی فقط شمارش و درصد است. محاسبه در turnover_metrics است.
"""
from __future__ import annotations

import asyncio
import io
import logging
import time
from collections import Counter
from datetime import datetime, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.org_tree import normalize_code
from app.core.persian_date import get_current_jalali_date
from app.models.employee import EmployeeMapping
from app.models.site import Site, SiteConnection
from app.models.termination_reason import REASON_GROUPS, TerminationReasonAlias, TerminationReasonCategory
from app.services.turnover_metrics import (
    GROUP_LABELS,
    RETENTION_POINTS,
    Record,
    compute_report,
    normalize_reason,
    parse_jalali_int,
    parse_month,
)

logger = logging.getLogger("faipco.turnover_report")

_CACHE_TTL_SECONDS = 300


class TurnoverReportError(Exception):
    """خطای قابل نمایش به کاربر (نگاشت ناقص، اتصال ناموفق، ...)."""


# ---------- خواندن از منبع و Cache ----------

_site_cache: dict[int, tuple[float, dict]] = {}


def clear_cache(site_id: int | None = None) -> None:
    """Cache داده‌ی خام یک سایت (یا همه) را پاک می‌کند؛ بعد از تغییر نگاشت."""
    if site_id is None:
        _site_cache.clear()
    else:
        _site_cache.pop(site_id, None)


class TurnoverReportService:
    """خواندن داده از منبع (با Cache پنج‌دقیقه‌ای)، مدیریت دسته‌ها و ساخت گزارش/Excel."""

    def __init__(self, db: AsyncSession):
        self.db = db

    # --- سایت‌ها ---
    async def configured_sites(self, allowed: set[int] | None) -> list[dict]:
        """سایت‌های مجاز که نگاشت تاریخ استخدام و ترک کار دارند: [{id, name, start_month}]."""
        result = await self.db.execute(
            select(Site, EmployeeMapping).join(EmployeeMapping, EmployeeMapping.site_id == Site.id).order_by(Site.name)
        )
        sites = []
        for site, mapping in result.all():
            if allowed is not None and site.id not in allowed:
                continue
            if not (mapping.hire_date_column or "").strip() or not (mapping.termination_date_column or "").strip():
                continue
            sites.append({"id": site.id, "name": site.name, "start_month": mapping.turnover_start_month})
        return sites

    async def load_site(self, site_id: int, refresh: bool = False) -> dict:
        """
        ردیف‌های پرسنل سایت را (فقط ستون‌های غیرشناسایی‌کننده) از منبع می‌خواند و بعد از فیلتر شعبه
        و درخت واحدها برمی‌گرداند: {"records", "dept_names", "position_names", "education_names"}.
        """
        cached = _site_cache.get(site_id)
        if cached and not refresh and time.monotonic() - cached[0] < _CACHE_TTL_SECONDS:
            return cached[1]

        from app.sync_engine.sync_service import SyncError, SyncService  # جلوگیری از import حلقه‌ای

        sync = SyncService(self.db)
        try:
            conn: SiteConnection = await sync._get_site_connection(site_id)
            mapping: EmployeeMapping = await sync._get_mapping(site_id)
        except SyncError as e:
            raise TurnoverReportError(str(e)) from e
        hire_col = (mapping.hire_date_column or "").strip()
        term_col = (mapping.termination_date_column or "").strip()
        if not hire_col or not term_col:
            raise TurnoverReportError("در نگاشت پرسنل این سایت، ستون تاریخ استخدام و ستون تاریخ ترک کار باید تنظیم شوند.")

        fields = {
            "hire": hire_col,
            "term": term_col,
            "active": (mapping.is_active_column or "").strip(),
            "reason": (mapping.termination_reason_column or "").strip(),
            "dept": (mapping.department_column or "").strip(),
            "gender": (mapping.gender_column or "").strip(),
            "birth": (mapping.birth_date_column or "").strip(),
            "position": (mapping.position_column or "").strip(),
            "education": (mapping.education_column or "").strip(),
            "branch": (mapping.branch_code_column or "").strip(),
        }
        columns = list(dict.fromkeys(v for v in fields.values() if v))
        adapter = sync._build_adapter(conn)
        try:
            rows = await adapter.fetch_rows(mapping.table_name, columns)
            rows = sync._filter_branch(mapping, rows)
            org_scope = await sync._resolve_org_scope(site_id, conn, mapping, adapter)
            dept_names = await sync._load_lookup_table(
                adapter, mapping.department_lookup_table, mapping.department_lookup_id_column, mapping.department_lookup_name_column
            )
            position_names = await sync._load_lookup_table(
                adapter, mapping.position_lookup_table, mapping.position_lookup_id_column, mapping.position_lookup_name_column
            )
            education_names = await sync._load_lookup_table(
                adapter, mapping.education_lookup_table, mapping.education_lookup_id_column, mapping.education_lookup_name_column
            )
        except SyncError as e:
            raise TurnoverReportError(str(e)) from e
        except Exception as e:  # noqa: BLE001 - خطای اتصال/کوئری منبع با پیام قابل فهم
            logger.exception("خواندن داده‌ی گزارش جذب و ترک کار (سایت %s) ناموفق بود", site_id)
            raise TurnoverReportError("خواندن اطلاعات از دیتابیس منبع ناموفق بود؛ اتصال و نگاشت پرسنل را بررسی کنید.") from e

        unassigned = 0
        if org_scope is not None:
            assignment, _siblings = org_scope
            own = []
            for r in rows:
                owner = assignment.get(normalize_code(r.get(fields["dept"])))
                if owner == site_id:
                    own.append(r)
                elif owner is None:
                    unassigned += 1  # واحدش زیر هیچ ریشه‌ای نیست (یا دیگر وجود ندارد)؛ در هشدار گزارش شمرده می‌شود
            rows = own

        def _code(row, key):
            col = fields[key]
            return normalize_code(row.get(col)) if col else None

        records = []
        for row in rows:
            term = parse_jalali_int(row.get(term_col))
            if fields["active"]:
                raw_active = row.get(fields["active"])
                active = sync._coerce_is_active(raw_active)
                if mapping.is_active_inverted and raw_active is not None:  # NULL در ستون برعکس (IsCut) = فعال
                    active = not active
                left = not active
            else:
                left = term is not None
            reason_raw = str(row.get(fields["reason"]) or "") if fields["reason"] else ""
            records.append(
                Record(
                    hire=parse_jalali_int(row.get(hire_col)),
                    term=term if left else None,
                    left=left,
                    reason_raw=reason_raw.strip(),
                    reason_norm=normalize_reason(reason_raw),
                    dept=_code(row, "dept"),
                    gender=sync._normalize_gender(row.get(fields["gender"])) if fields["gender"] else None,
                    birth=parse_jalali_int(row.get(fields["birth"])) if fields["birth"] else None,
                    position=_code(row, "position"),
                    education=_code(row, "education"),
                )
            )
        data = {
            "unassigned": unassigned,
            "conn_key": (conn.db_type, (conn.host or "").strip().lower(), conn.port, (conn.database_name or "").strip().lower()),
            "records": records,
            "dept_names": dept_names,
            "position_names": position_names,
            "education_names": education_names,
        }
        _site_cache[site_id] = (time.monotonic(), data)
        return data

    # --- دسته‌ها ---
    async def list_categories(self) -> list[dict]:
        result = await self.db.execute(
            select(TerminationReasonCategory).order_by(TerminationReasonCategory.sort_order, TerminationReasonCategory.id)
        )
        return [
            {"id": c.id, "key": c.key, "title": c.title, "group": c.group, "legal_basis": c.legal_basis, "sort_order": c.sort_order}
            for c in result.scalars().all()
        ]

    async def alias_map(self, seen: dict[str, str]) -> dict[str, int | None]:
        """
        نگاشت «متن نرمال‌شده → شناسه دسته» را برمی‌گرداند و متن‌های تازه‌ی دیده‌شده (seen: نرمال → نمونه‌ی خام)
        را با دسته‌ی خالی ذخیره می‌کند تا در پنل دسته‌بندی شوند.
        """
        result = await self.db.execute(select(TerminationReasonAlias))
        aliases = {a.normalized_text: a.category_id for a in result.scalars().all()}
        new = {norm: raw for norm, raw in seen.items() if norm not in aliases}
        if new:
            for norm, raw in new.items():
                self.db.add(TerminationReasonAlias(normalized_text=norm, sample_text=(raw or "(بدون علت)")[:400], category_id=None))
                aliases[norm] = None
            try:
                await self.db.commit()
            except Exception:  # noqa: BLE001 - درج هم‌زمان توسط درخواست دیگر؛ گزارش نباید خراب شود
                await self.db.rollback()
        return aliases

    async def list_aliases(self, counts: Counter | None) -> list[dict]:
        result = await self.db.execute(select(TerminationReasonAlias).order_by(TerminationReasonAlias.normalized_text))
        rows = [
            {
                "id": a.id,
                "text": a.sample_text,
                "normalized_text": a.normalized_text,
                "category_id": a.category_id,
                "count": counts.get(a.normalized_text, 0) if counts is not None else None,
            }
            for a in result.scalars().all()
        ]
        rows.sort(key=lambda r: (r["category_id"] is not None, -(r["count"] or 0), r["text"]))
        return rows

    # --- گزارش ---
    async def build(
        self,
        site_ids: list[int],
        start_months: list[str | None],
        from_month: str | None,
        to_month: str | None,
        dept: str | None,
        gender: int | None,
        refresh: bool = False,
    ) -> dict:
        """داده‌ی سایت‌های انتخابی را ترکیب و گزارش را محاسبه می‌کند."""
        records: list[Record] = []
        dept_names: dict[str, str] = {}
        position_names: dict[str, str] = {}
        education_names: dict[str, str] = {}
        loaded = [(sid, await self.load_site(sid, refresh=refresh)) for sid in site_ids]
        # سایت‌هایی با دیتابیس منبع متفاوت ممکن است کد واحد یکسان داشته باشند؛ آن‌وقت کد واحد با شناسه‌ی سایت یکتا می‌شود
        prefix = len({d["conn_key"] for _sid, d in loaded}) > 1
        unassigned_by_source: dict = {}
        for sid, data in loaded:
            # سایت‌های هم‌منبع همان افراد بی‌سایت را می‌بینند؛ برای هر منبع یک‌بار شمرده می‌شود
            key = data["conn_key"]
            unassigned_by_source[key] = max(unassigned_by_source.get(key, 0), data.get("unassigned", 0))
            for r in data["records"]:
                copy = Record(**vars(r))  # کپی؛ prepare_records تاریخ را اصلاح می‌کند
                if prefix and copy.dept:
                    copy.dept = f"{sid}:{copy.dept}"
                records.append(copy)
            names = data["dept_names"]
            dept_names.update({f"{sid}:{k}": v for k, v in names.items()} if prefix else names)
            position_names.update(data["position_names"])
            education_names.update(data["education_names"])
        seen = {r.reason_norm: r.reason_raw for r in records if r.left}
        alias_category = await self.alias_map(seen)
        categories = await self.list_categories()
        starts = [parse_month(s) for s in start_months if parse_month(s) is not None]
        y, m, d = get_current_jalali_date()
        report = await asyncio.to_thread(
            compute_report,
            records,
            alias_category=alias_category,
            categories=categories,
            start_idx=max(starts) if starts else None,
            from_idx=parse_month(from_month),
            to_idx=parse_month(to_month),
            today=y * 10000 + m * 100 + d,
            dept_names=dept_names,
            position_names=position_names,
            education_names=education_names,
            dept_filter=(dept or "").strip() or None,
            gender_filter=gender,
        )
        unassigned = sum(unassigned_by_source.values())
        if unassigned:
            report["warnings"]["unassigned_unit"] = unassigned
        report["categories"] = categories
        report["generated_at"] = datetime.now(timezone.utc).isoformat()
        return report

    async def reason_counts(self, site_ids: list[int]) -> Counter:
        """تعداد هر متن نرمال‌شده‌ی علت در بین ترک‌کرده‌های سایت‌ها (برای صفحه‌ی دسته‌بندی)."""
        counts: Counter = Counter()
        for sid in site_ids:
            data = await self.load_site(sid)
            counts.update(r.reason_norm for r in data["records"] if r.left)
        # متن‌های تازه (هنوز در جدول نیامده) هم ثبت می‌شوند تا در فهرست دسته‌بندی دیده شوند
        seen = {}
        for sid in site_ids:
            for r in (await self.load_site(sid))["records"]:
                if r.left:
                    seen.setdefault(r.reason_norm, r.reason_raw)
        await self.alias_map(seen)
        return counts


# ---------- Excel ----------

def build_xlsx(report: dict, title: str) -> bytes:
    """خروجی Excel چندبرگه‌ای گزارش (راست‌به‌چپ)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="185E95")

    def sheet(name: str, headers: list[str], rows: list[list], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        for i, _h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
        ws.freeze_panes = "A2"
        return ws

    k = report["kpis"]
    r = report["range"]
    kpi_rows = [
        ["عنوان گزارش", title],
        ["بازه", f"{r['from_month']} تا {r['to_month']}"],
        ["پرسنل ابتدای دوره", k["start_headcount"]],
        ["پرسنل انتهای دوره", k["end_headcount"]],
        ["میانگین پرسنل", k["avg_headcount"]],
        ["استخدام", k["hires"]],
        ["ترک کار", k["separations"]],
        ["خالص", k["net"]],
        ["نرخ ترک کار دوره (%)", k["turnover_rate_period"]],
        ["نرخ ترک کار سالانه‌شده (%)", k["turnover_rate_annualized"]],
        ["نرخ ترک به خواست کارگر، سالانه‌شده (%)", k["voluntary_rate_annualized"]],
        ["نرخ ترک به خواست کارفرما، سالانه‌شده (%)", k["involuntary_rate_annualized"]],
        ["نرخ جذب سالانه‌شده (%)", k["hire_rate_annualized"]],
        ["نسبت جایگزینی (استخدام ÷ ترک)", k["replacement_ratio"]],
        ["سهم ترک در ۳۰ روز اول (%)", k["exit_within_30_share"]],
        ["سهم ترک در ۹۰ روز اول (%)", k["exit_within_90_share"]],
        ["سهم ترک در سال اول (%)", k["exit_within_365_share"]],
        ["ریزش نیروی جدید در ۹۰ روز (%)", k["new_hire_90d_attrition"]],
        ["ماندگاری سال اول (%)", k["first_year_retention"]],
        ["میانگین مدت خدمت هنگام ترک (ماه)", k["avg_tenure_at_exit_months"]],
    ]
    sheet("خلاصه", ["شاخص", "مقدار"], kpi_rows, first=True)

    groups = ["voluntary", "involuntary", "probation", "other", "uncategorized"]
    sheet(
        "ماهانه",
        ["ماه", "پرسنل اول ماه", "پرسنل آخر ماه", "میانگین پرسنل", "استخدام", "ترک کار", "خالص"]
        + [GROUP_LABELS[g] for g in groups]
        + ["نرخ ترک (%)", "نرخ ترک به خواست کارگر (%)", "نرخ جذب (%)", "نرخ متحرک ۱۲ماهه (%)"],
        [
            [m["month"], m["start_headcount"], m["end_headcount"], m["avg_headcount"], m["hires"], m["separations"], m["net"]]
            + [m["by_group"][g] for g in groups]
            + [m["turnover_rate"], m["voluntary_rate"], m["hire_rate"], m["rolling12_rate"]]
            for m in report["monthly"]
        ],
    )
    sheet(
        "سالانه",
        ["سال", "تعداد ماه", "پرسنل ابتدا", "پرسنل انتها", "استخدام", "ترک کار", "خالص", "نرخ ترک (%)", "نرخ سالانه‌شده (%)"],
        [[y["year"], y["months"], y["start_headcount"], y["end_headcount"], y["hires"], y["separations"], y["net"],
          y["turnover_rate"], y["turnover_rate_annualized"]] for y in report["yearly"]],
    )
    sheet(
        "علت ترک کار",
        ["دسته", "گروه", "مبنای قانونی", "تعداد", "سهم (%)", "متن‌های ثبت‌شده"],
        [[x["title"], GROUP_LABELS.get(x["group"], x["group"]), x.get("legal_basis") or "", x["count"], x["share"],
          "، ".join(f"{t['text']} ({t['count']})" for t in x["texts"])] for x in report["reasons"]],
    )
    sheet("مدت خدمت", ["مدت خدمت هنگام ترک", "تعداد", "سهم (%)"],
          [[b["label"], b["count"], b["share"]] for b in report["tenure_buckets"]])
    sheet(
        "ماندگاری",
        ["گروه استخدام", "تعداد"] + [f"ماندگاری {d} روز (%)" for d in RETENTION_POINTS],
        [[c["label"], c["size"]] + [p["retention"] for p in c["points"]] for c in report["cohorts"]],
    )
    sheet(
        "واحدها",
        ["واحد", "استخدام", "ترک کار", "به خواست کارگر", "به خواست کارفرما", "میانگین پرسنل", "پرسنل انتهای دوره", "نرخ ترک سالانه‌شده (%)"],
        [[d["name"], d["hires"], d["separations"], d["voluntary"], d["involuntary"], d["avg_headcount"], d["end_headcount"],
          d["turnover_rate_annualized"]] for d in report["departments"]],
    )
    demo_rows = []
    for key, label in (("gender", "جنسیت"), ("age", "گروه سنی"), ("education", "مدرک تحصیلی"), ("positions", "سمت")):
        for x in report["demographics"][key]:
            demo_rows.append([label, x["label"], x["hires"], x["separations"]])
    sheet("ترکیب جمعیتی", ["بعد", "مقدار", "استخدام", "ترک کار"], demo_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "REASON_GROUPS",
    "GROUP_LABELS",
    "Record",
    "TurnoverReportError",
    "TurnoverReportService",
    "build_xlsx",
    "clear_cache",
    "compute_report",
    "normalize_reason",
    "parse_jalali_int",
]
