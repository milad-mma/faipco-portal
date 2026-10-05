"""
سرویس ماژول «مشخصات خانوادگی» (پایه‌ی حق تاهل و حق اولاد).

- تنظیمات (فیلدها، مدارک، قواعد شمول، هشدارها) در system_settings با کلید family_settings
- فرم کارمند: ثبت/ویرایش پرونده، آپلود/حذف مدرک
- پنل منابع انسانی: فهرست با شمول محاسبه‌شده، جزئیات، تأیید/رد/بازگشت برای ویرایش، سابقه بیمه، خروجی Excel
- پاک‌سازی مدارکی که آپلود شده ولی هرگز در فرم ثبت نشده‌اند

قواعد و اعتبارسنجی در core/family_rules است؛ این سرویس فقط دیتابیس و جریان کار را مدیریت می‌کند.
مستقل از ماژول بیمه تکمیلی است.
"""
from __future__ import annotations

import asyncio
import io
import json
import logging
from datetime import datetime, timedelta, timezone

from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload, undefer

from app.core import family_rules as rules
from app.core.document_sanitizer import DocumentRejected, sanitize_document, sanitize_file_name, sniff_content_type
from app.core.persian_date import get_current_jalali_date
from app.core.text_normalize import normalize_search_text
from app.models.employee import Department, Employee
from app.models.family import FamilyChangeLog, FamilyDocument, FamilyMember, FamilyProfile
from app.models.notice import Notice, NoticePriority, NoticeStatus, NoticeTarget, NoticeTargetType, NoticeType
from app.models.site import Site
from app.models.system_setting import SystemSetting
from app.models.user import User

logger = logging.getLogger(__name__)

SETTINGS_KEY = "family_settings"
DOCUMENT_MAX_BYTES = 10 * 1024 * 1024
DOCUMENT_ALLOWED_TYPES = {
    "image/jpeg", "image/png", "image/gif", "image/webp", "image/bmp", "image/tiff", "application/pdf",
}  # fmt: skip
MAX_UNLINKED_DOCUMENTS = 15
UNLINKED_MAX_AGE_HOURS = 72

ACTION_LABELS = {
    "submitted": "ثبت/ویرایش توسط پرسنل",
    "approved": "تأیید",
    "rejected": "رد",
    "returned": "بازگشت برای ویرایش",
    "hr_fields": "تغییر سابقه بیمه / یادداشت",
}


class FamilyError(Exception):
    """خطای قابل نمایش به کاربر (پیام فارسی)."""


class FamilyForbiddenError(FamilyError):
    """ماژول غیرفعال یا پرونده قفل است."""


def today_jalali() -> tuple[int, int, int]:
    return get_current_jalali_date()


def since_hire_days(employee: Employee, as_of: tuple[int, int, int]) -> int | None:
    """روزهای سپری‌شده از تاریخ استخدام (از کاراوب) تا as_of؛ تاریخ نامعتبر/خالی → None."""
    hire = rules.parse_jalali(getattr(employee, "hire_date_jalali", None))
    if hire is None or hire > as_of:
        return None
    return rules.days_between(hire, as_of)


def insurance_info(
    settings: dict, employee: Employee, profile: FamilyProfile | None, data: dict | None, as_of
) -> dict:
    """
    سابقه‌ی بیمه‌ی یک پرسنل در تاریخ as_of با منبع و اجزا.
    data: نسخه‌ی پرونده (تأییدشده یا جاری) برای سابقه‌ی پیش از استخدامِ اعلام‌شده توسط پرسنل.
    """
    since = since_hire_days(employee, as_of)
    hr_prior = profile.hr_prior_insurance_days if profile else None
    employee_prior = (data or {}).get("prior_insurance_days")
    prior = hr_prior if hr_prior is not None else employee_prior
    override = profile.insurance_days if profile else None
    days, source = rules.total_insurance_days(settings, since, prior, override, prior_by_hr=hr_prior is not None)
    return {
        "days": days,
        "source": source,
        "source_label": rules.INSURANCE_SOURCES.get(source) if source else None,
        "since_hire_days": since,
        "employee_prior_days": employee_prior,
        "hr_prior_days": hr_prior,
        "hr_total_days": override,
    }


class FamilyService:
    def __init__(self, db: AsyncSession):
        self.db = db

    # ---------- تنظیمات ----------

    async def get_settings(self) -> dict:
        row = await self.db.get(SystemSetting, SETTINGS_KEY)
        stored = None
        if row is not None:
            try:
                stored = json.loads(row.value)
            except (ValueError, TypeError):
                stored = None
        return rules.sanitize_settings(stored)

    async def update_settings(self, patch: dict) -> dict:
        current = await self.get_settings()
        new = rules.merge_settings(current, patch)
        row = await self.db.get(SystemSetting, SETTINGS_KEY)
        value = json.dumps(new, ensure_ascii=False)
        if row is None:
            self.db.add(SystemSetting(key=SETTINGS_KEY, value=value))
        else:
            row.value = value
        await self.db.commit()
        return new

    @staticmethod
    def settings_meta() -> dict:
        """تعریف فیلدها/مدارک/گزینه‌ها برای ساخت صفحه‌ی تنظیمات و فرم در فرانت."""
        return {
            "field_defs": rules.FIELD_DEFS,
            "doc_types": rules.DOC_TYPES,
            "marital_statuses": rules.MARITAL_STATUSES,
            "relations": rules.RELATIONS,
            "custody_options": rules.CUSTODY_OPTIONS,
            "statuses": rules.PROFILE_STATUSES,
            "marriage_female_modes": rules.MARRIAGE_FEMALE_MODES,
            "child_female_modes": rules.CHILD_FEMALE_MODES,
        }

    def form_for_employee(self, settings: dict) -> dict:
        """تنظیمات لازم برای فرم کارمند (بدون قواعد شمول، که فقط به HR مربوط است)."""
        return {
            "fields": settings["fields"],
            "documents": {k: {"mode": v["mode"]} for k, v in settings["documents"].items()},
            "notes": settings["notes"],
            # سؤال‌های تحصیل و گواهی تحصیل پسر/دختر از این سن به بعد پرسیده می‌شوند
            "son_study_age": settings["rules"]["child"]["son_study_age"],
            "daughter_study_age": settings["rules"]["child"]["daughter_study_age"],
            **self.settings_meta(),
        }

    # ---------- خواندن پرونده ----------

    async def get_profile_by_employee(self, employee_id: int) -> FamilyProfile | None:
        result = await self.db.execute(
            select(FamilyProfile)
            .options(
                selectinload(FamilyProfile.members).selectinload(FamilyMember.documents),
                selectinload(FamilyProfile.documents),
            )
            .where(FamilyProfile.employee_id == employee_id)
            .execution_options(populate_existing=True)
        )
        return result.scalar_one_or_none()

    async def get_profile(self, profile_id: int, site_ids: set[int] | None) -> tuple[FamilyProfile, Employee] | None:
        result = await self.db.execute(
            select(FamilyProfile, Employee)
            .options(
                selectinload(FamilyProfile.members).selectinload(FamilyMember.documents),
                selectinload(FamilyProfile.documents),
            )
            .join(Employee, Employee.id == FamilyProfile.employee_id)
            .where(FamilyProfile.id == profile_id)
            .execution_options(populate_existing=True)
        )
        row = result.first()
        if row is None:
            return None
        profile, employee = row
        if site_ids is not None and employee.site_id not in site_ids:
            return None
        return profile, employee

    @staticmethod
    def _doc_out(doc: FamilyDocument, settings: dict, today: tuple[int, int, int], member: dict | None = None) -> dict:
        expiry = rules.document_expiry(settings, {"doc_type": doc.doc_type, "uploaded": doc.uploaded_jalali}, member)
        return {
            "id": doc.id,
            "doc_type": doc.doc_type,
            "file_name": doc.file_name,
            "content_type": doc.content_type,
            "size_bytes": doc.size_bytes,
            "uploaded_at": doc.uploaded_at,
            "uploaded_jalali": doc.uploaded_jalali,
            "expires_on": rules.format_jalali(expiry),
            "expired": bool(expiry and expiry < today),
        }

    def profile_out(self, profile: FamilyProfile, settings: dict) -> dict:
        today = today_jalali()
        members = []
        for m in profile.members:
            item = {col: getattr(m, col) for col in rules.MEMBER_COLUMNS}
            item.update(
                id=m.id,
                member_type=m.member_type,
                first_name=m.first_name,
                last_name=m.last_name,
                documents=[
                    self._doc_out(d, settings, today, item) for d in m.documents if d.linked and not d.archived
                ],
            )
            members.append(item)
        return {
            "id": profile.id,
            "status": profile.status,
            "marital_status": profile.marital_status,
            "marriage_date": profile.marriage_date,
            "separation_date": profile.separation_date,
            "is_head_of_household": profile.is_head_of_household,
            "has_children": profile.has_children,
            "prior_insurance_days": profile.prior_insurance_days,
            "submitted_at": profile.submitted_at,
            "reviewed_at": profile.reviewed_at,
            "review_note": profile.review_note,
            "effective_date": profile.effective_date,
            "approved_at": profile.approved_at,
            "documents": [
                self._doc_out(d, settings, today)
                for d in profile.documents
                if d.linked and not d.archived and d.member_id is None
            ],
            "unlinked_documents": [self._doc_out(d, settings, today) for d in profile.documents if not d.linked],
            "members": members,
        }

    @staticmethod
    def snapshot(profile: FamilyProfile) -> dict:
        """داده‌ی پرونده به قالب family_rules (برای محاسبه‌ی شمول و نسخه‌ی تأییدشده)."""

        def docs_meta(docs):
            return [
                {"id": d.id, "doc_type": d.doc_type, "uploaded": d.uploaded_jalali, "file_name": d.file_name}
                for d in docs
                if d.linked and not d.archived
            ]

        return {
            "marital_status": profile.marital_status,
            "marriage_date": profile.marriage_date,
            "separation_date": profile.separation_date,
            "is_head_of_household": profile.is_head_of_household,
            "has_children": profile.has_children,
            "prior_insurance_days": profile.prior_insurance_days,
            "docs": docs_meta([d for d in profile.documents if d.member_id is None]),
            "members": [
                {
                    "member_type": m.member_type,
                    "first_name": m.first_name,
                    "last_name": m.last_name,
                    **{col: getattr(m, col) for col in rules.MEMBER_COLUMNS if col not in ("first_name", "last_name")},
                    "docs": docs_meta(m.documents),
                }
                for m in profile.members
            ],
        }

    # ---------- کارمند ----------

    @staticmethod
    def can_edit(profile: FamilyProfile | None, settings: dict) -> tuple[bool, str | None]:
        if not settings["enabled"]:
            return False, "ثبت مشخصات خانوادگی در حال حاضر غیرفعال است."
        if profile is not None and profile.status == "approved" and settings["lock_after_approval"]:
            return False, "مشخصات شما تأیید شده است؛ برای هر تغییر به واحد منابع انسانی مراجعه کنید."
        return True, None

    async def my_status(self, employee: Employee | None) -> dict:
        settings = await self.get_settings()
        profile = await self.get_profile_by_employee(employee.id) if employee else None
        can_edit, reason = self.can_edit(profile, settings) if employee else (False, None)
        return {
            "enabled": settings["enabled"],
            "can_edit": can_edit,
            "lock_reason": reason,
            "employee": {
                "first_name": employee.first_name,
                "last_name": employee.last_name,
                "personnel_code": employee.personnel_code,
                "gender": employee.gender,
            }
            if employee
            else None,
            "profile": self.profile_out(profile, settings) if profile else None,
            "form": self.form_for_employee(settings),
            # آیا سؤال «سابقه‌ی بیمه‌ی پیش از استخدام» پرسیده شود (با/بدون فرزند)؛ اگر منابع انسانی خودش سابقه را
            # ثبت کرده باشد پرسیده نمی‌شود
            "prior_insurance": self._prior_flags(employee, profile, settings) if employee else None,
        }

    @staticmethod
    def _prior_flags(employee: Employee, profile: FamilyProfile | None, settings: dict) -> dict:
        hr_set = profile is not None and (profile.insurance_days is not None or profile.hr_prior_insurance_days is not None)
        since = since_hire_days(employee, today_jalali())
        return {
            "with_children": not hr_set and rules.prior_insurance_needed(settings, since, True),
            "without_children": not hr_set and rules.prior_insurance_needed(settings, since, False),
        }

    async def _get_or_create_shell(self, employee: Employee) -> FamilyProfile:
        profile = await self.get_profile_by_employee(employee.id)
        if profile is not None:
            return profile
        profile = FamilyProfile(employee_id=employee.id, status="draft")
        self.db.add(profile)
        try:
            await self.db.flush()
        except IntegrityError:
            # درخواست هم‌زمان دیگری (مثلاً دو آپلود اول پشت‌سرهم) پرونده را ساخته است
            await self.db.rollback()
            existing = await self.get_profile_by_employee(employee.id)
            if existing is None:
                raise
            return existing
        await self.db.refresh(profile, attribute_names=["members", "documents"])
        return profile

    async def upload_document(self, employee: Employee, doc_type: str, file_name: str, content: bytes) -> dict:
        settings = await self.get_settings()
        if doc_type not in rules.DOC_TYPES or settings["documents"][doc_type]["mode"] == "hidden":
            raise FamilyError("نوع مدرک نامعتبر است.")
        profile = await self.get_profile_by_employee(employee.id)
        ok, reason = self.can_edit(profile, settings)
        if not ok:
            raise FamilyForbiddenError(reason)
        if len(content) > DOCUMENT_MAX_BYTES:
            raise FamilyError("حجم فایل نباید بیشتر از ۱۰ مگابایت باشد.")
        detected = sniff_content_type(content)
        if detected not in DOCUMENT_ALLOWED_TYPES:
            raise FamilyError("فقط فایل تصویری (JPG/PNG/GIF/WEBP/BMP/TIFF) یا PDF پذیرفته می‌شود.")
        try:
            content, detected = await asyncio.to_thread(sanitize_document, content, detected)
        except DocumentRejected as e:
            raise FamilyError(str(e)) from e
        if len(content) > DOCUMENT_MAX_BYTES:
            raise FamilyError("حجم فایل نباید بیشتر از ۱۰ مگابایت باشد.")
        profile = profile or await self._get_or_create_shell(employee)
        if sum(1 for d in profile.documents if not d.linked) >= MAX_UNLINKED_DOCUMENTS:
            raise FamilyError(
                f"حداکثر {MAX_UNLINKED_DOCUMENTS} مدرک ثبت‌نشده مجاز است؛ لطفاً فرم را ثبت کنید یا مدارک اضافه را حذف کنید."
            )
        doc = FamilyDocument(
            profile_id=profile.id,
            doc_type=doc_type,
            linked=False,
            file_name=sanitize_file_name(file_name, detected),
            content_type=detected,
            size_bytes=len(content),
            data=content,
            uploaded_jalali=rules.format_jalali(today_jalali()),
        )
        self.db.add(doc)
        await self.db.commit()
        await self.db.refresh(doc)
        return self._doc_out(doc, settings, today_jalali())

    async def get_document_for_download(
        self,
        document_id: int,
        employee_id: int | None,
        allowed_site_ids: set[int] | None = frozenset(),
        with_data: bool = True,
    ) -> FamilyDocument | None:
        """
        with_data=True (پیش‌فرض): ستون deferred «data» هم بارگذاری می‌شود (برای دانلود). برای کارهایی که
        فقط فراداده لازم دارند (مثل حذف) False بدهید تا فایل حجیم بی‌دلیل خوانده نشود.
        """
        stmt = select(FamilyDocument, FamilyProfile.employee_id, Employee.site_id)
        if with_data:
            stmt = stmt.options(undefer(FamilyDocument.data))
        result = await self.db.execute(
            stmt.join(FamilyProfile, FamilyProfile.id == FamilyDocument.profile_id)
            .join(Employee, Employee.id == FamilyProfile.employee_id)
            .where(FamilyDocument.id == document_id)
        )
        row = result.first()
        if row is None:
            return None
        doc, owner_id, site_id = row
        if employee_id is not None and owner_id == employee_id:
            return doc
        if allowed_site_ids is None or site_id in allowed_site_ids:
            return doc
        return None

    async def delete_own_document(self, document_id: int, employee: Employee) -> bool:
        settings = await self.get_settings()
        profile = await self.get_profile_by_employee(employee.id)
        ok, reason = self.can_edit(profile, settings)
        if not ok:
            raise FamilyForbiddenError(reason)
        doc = await self.get_document_for_download(document_id, employee.id, with_data=False)
        if doc is None:
            return False
        if doc.linked:
            # مدرک ثبت‌شده فقط با ویرایش و ثبت دوباره‌ی فرم حذف می‌شود (تا پرونده‌ی ثبت‌شده ناقص نشود)
            raise FamilyError("این مدرک در فرم ثبت‌شده است؛ آن را در فرم حذف و فرم را دوباره ثبت کنید.")
        await self.db.delete(doc)
        await self.db.commit()
        return True

    async def save(self, employee: Employee, payload: dict, user: User) -> dict:
        """
        فرم کارمند را اعتبارسنجی و به‌طور کامل جایگزین می‌کند و وضعیت را «در انتظار بررسی» می‌گذارد.
        نسخه‌ی تأییدشده‌ی قبلی (approved_data) تا تأیید دوباره در گزارش استفاده می‌شود.
        """
        settings = await self.get_settings()
        profile = await self.get_profile_by_employee(employee.id)
        ok, reason = self.can_edit(profile, settings)
        if not ok:
            raise FamilyForbiddenError(reason)
        today = today_jalali()

        # ترتیب اعضا مثل خروجی validate_payload: همسر اول، سپس فرزندان به ترتیب فرم
        members_in = payload.get("members") or []
        ordered = [m for m in members_in if m.get("member_type") == "spouse"] + [
            m for m in members_in if m.get("member_type") != "spouse"
        ]
        try:
            flags = self._prior_flags(employee, profile, settings)
            normalized = rules.validate_payload(
                {**payload, "members": ordered},
                settings,
                today,
                since_hire_days(employee, today),
                ask_prior=flags["with_children"] or flags["without_children"],
            )
        except rules.FamilyRuleError as e:
            raise FamilyError(str(e)) from e

        profile = profile or await self._get_or_create_shell(employee)
        own_docs = {d.id: d for d in profile.documents}

        # بررسی مدارک ارجاع‌شده: متعلق به همین پرونده، بدون تکرار، و نوعش برای همان جایگاه مجاز باشد
        used: set[int] = set()

        def take(ids: list[int], allowed: list[str], who: str) -> list[FamilyDocument]:
            docs = []
            for doc_id in ids or []:
                doc = own_docs.get(doc_id)
                if doc is None or doc_id in used:
                    raise FamilyError("مدرک ارجاع‌شده یافت نشد یا متعلق به شما نیست.")
                if doc.doc_type not in allowed or doc.archived:
                    # مدرکی که با پاسخ‌های جدید موضوعیت ندارد (مثلاً گواهی تحصیل پس از «خیر») کنار گذاشته می‌شود
                    continue
                used.add(doc_id)
                docs.append(doc)
            return docs

        profile_docs = take(payload.get("document_ids") or [], rules.allowed_documents(settings, normalized, None, today), "پرونده")
        member_docs: list[list[FamilyDocument]] = []
        for spec, raw in zip(normalized["members"], ordered):
            who = rules.MEMBER_TYPES[spec["member_type"]] + f" «{spec['first_name']}»"
            member_docs.append(take(raw.get("document_ids") or [], rules.allowed_documents(settings, normalized, spec, today), who))

        # مدارک اجباری
        check = {
            **normalized,
            "docs": [{"doc_type": d.doc_type} for d in profile_docs],
            "members": [{**spec, "docs": [{"doc_type": d.doc_type} for d in docs]} for spec, docs in zip(normalized["members"], member_docs)],
        }
        missing = rules.missing_documents(settings, check, today)
        if missing:
            raise FamilyError("مدارک زیر را پیوست کنید: " + "، ".join(missing))

        # جایگزینی اعضا
        old_member_ids = [m.id for m in profile.members]
        for idx, (spec, docs) in enumerate(zip(normalized["members"], member_docs)):
            member = FamilyMember(profile_id=profile.id, sort_order=idx, **spec)
            self.db.add(member)
            await self.db.flush()
            for d in docs:
                d.member_id = member.id
                d.linked = True
        for d in profile_docs:
            d.member_id = None
            d.linked = True
        await self.db.flush()
        if old_member_ids:
            await self.db.execute(delete(FamilyMember).where(FamilyMember.id.in_(old_member_ids)))
        # مدارکی که در فرم جدید ارجاع نشده‌اند (کارمند حذفشان کرده) پاک می‌شوند؛ مگر جزو نسخه‌ی تأییدشده باشند
        # که تا تأیید بعدی به‌صورت بایگانی می‌مانند (گزارش حقوق هنوز به آن نسخه تکیه دارد)
        approved_ids = _approved_doc_ids(profile.approved_data)
        stale = [doc_id for doc_id in own_docs if doc_id not in used and doc_id not in approved_ids]
        for doc_id, doc in own_docs.items():
            if doc_id not in used and doc_id in approved_ids:
                doc.archived = True
                doc.member_id = None
            elif doc_id in used:
                doc.archived = False
        if stale:
            await self.db.execute(delete(FamilyDocument).where(FamilyDocument.id.in_(stale)))

        profile.marital_status = normalized["marital_status"]
        profile.marriage_date = normalized["marriage_date"]
        profile.separation_date = normalized["separation_date"]
        profile.is_head_of_household = normalized["is_head_of_household"]
        profile.has_children = normalized["has_children"]
        profile.prior_insurance_days = normalized["prior_insurance_days"]
        profile.status = "pending"
        profile.submitted_at = datetime.now(timezone.utc)
        profile.review_note = None
        await self.db.flush()

        profile = await self.get_profile_by_employee(employee.id)
        self._log(profile, "submitted", user, snapshot=self.snapshot(profile))
        await self.db.commit()
        profile = await self.get_profile_by_employee(employee.id)
        return self.profile_out(profile, settings)

    # ---------- منابع انسانی ----------

    def _log(self, profile: FamilyProfile, action: str, user: User | None, note=None, effective_date=None, snapshot=None):
        self.db.add(
            FamilyChangeLog(
                profile_id=profile.id,
                action=action,
                actor_user_id=user.id if user else None,
                actor_name=(getattr(user, "full_name", None) or getattr(user, "username", None)) if user else None,
                note=note,
                effective_date=effective_date,
                snapshot=snapshot,
            )
        )

    async def _notify(self, employee_id: int, sender: User, title: str, body: str) -> int:
        notice = Notice(
            sender_id=sender.id,
            title=title[:255],
            body=body[:4000],
            priority=NoticePriority.high,
            status=NoticeStatus.published,
            notice_type=NoticeType.normal,
            publish_at=datetime.now(timezone.utc),
        )
        self.db.add(notice)
        await self.db.flush()
        self.db.add(NoticeTarget(notice_id=notice.id, target_type=NoticeTargetType.employee, target_id=employee_id))
        return notice.id

    async def list_profiles(
        self,
        site_ids: set[int] | None,
        search: str | None,
        status: str | None,
        flag: str | None,
        page: int,
        page_size: int,
        as_of: tuple[int, int, int],
    ) -> dict:
        """
        همه‌ی پرسنل فعال سایت‌های مجاز همراه پرونده (اگر باشد) و شمول محاسبه‌شده بر اساس نسخه‌ی تأییدشده.
        status: none / pending / approved / rejected / returned ؛ flag: marriage / children / warnings / no_insurance_days / pending_changes
        """
        settings = await self.get_settings()
        today = today_jalali()
        q = (
            select(Employee, Site.name, Department.name)
            .join(Site, Site.id == Employee.site_id)
            .outerjoin(Department, Department.id == Employee.department_id)
            .where(Employee.is_active.is_(True))
            .order_by(Employee.last_name, Employee.first_name)
        )
        if site_ids is not None:
            q = q.where(Employee.site_id.in_(site_ids))
        rows = (await self.db.execute(q)).all()
        emp_ids = [r[0].id for r in rows]
        profiles: dict[int, FamilyProfile] = {}
        if emp_ids:
            res = await self.db.execute(
                select(FamilyProfile)
                .options(
                    selectinload(FamilyProfile.members).selectinload(FamilyMember.documents),
                    selectinload(FamilyProfile.documents),
                )
                .where(FamilyProfile.employee_id.in_(emp_ids))
            )
            profiles = {p.employee_id: p for p in res.scalars().all()}

        term = normalize_search_text(search)
        items: list[dict] = []
        stats = {"employees": 0, "not_submitted": 0, "pending": 0, "approved": 0, "rejected": 0, "returned": 0,
                 "marriage_eligible": 0, "eligible_children": 0, "warnings": 0}  # fmt: skip
        for emp, site_name, dept_name in rows:
            item = self._list_item(emp, site_name, dept_name, profiles.get(emp.id), settings, as_of, today)
            stats["employees"] += 1
            stats["not_submitted" if item["status"] == "none" else item["status"]] += 1
            stats["marriage_eligible"] += 1 if item["marriage_eligible"] else 0
            stats["eligible_children"] += item["eligible_children"] or 0
            stats["warnings"] += 1 if item["warnings"] else 0
            if term and not any(
                term in (normalize_search_text(v) or "") for v in (emp.personnel_code, emp.first_name, emp.last_name, f"{emp.first_name} {emp.last_name}")
            ):
                continue
            if status and item["status"] != status:
                continue
            if flag == "marriage" and not item["marriage_eligible"]:
                continue
            if flag == "children" and not item["eligible_children"]:
                continue
            if flag == "warnings" and not item["warnings"]:
                continue
            if flag == "no_insurance_days" and item["insurance_days"] is not None:
                continue
            if flag == "pending_changes" and not item["has_pending_changes"]:
                continue
            items.append(item)
        # در انتظار بررسی‌ها اول
        items.sort(key=lambda i: 0 if i["status"] == "pending" else 1)
        total = len(items)
        start = (page - 1) * page_size
        return {"items": items[start : start + page_size], "total": total, "stats": stats}

    def _list_item(self, emp, site_name, dept_name, profile, settings, as_of, today) -> dict:
        status = "none" if profile is None or profile.status == "draft" else profile.status
        item = {
            "profile_id": profile.id if profile else None,
            "employee_id": emp.id,
            "personnel_code": emp.personnel_code,
            "first_name": emp.first_name,
            "last_name": emp.last_name,
            "site_id": emp.site_id,
            "site_name": site_name,
            "department_name": dept_name,
            "status": status,
            "marital_status": None,
            "sons": 0,
            "daughters": 0,
            "insurance_days": None,
            "insurance_source": None,
            "insurance_hr_total": profile.insurance_days if profile else None,
            "insurance_hr_prior": profile.hr_prior_insurance_days if profile else None,
            "marriage_eligible": None,
            "eligible_children": None,
            "has_pending_changes": bool(profile and profile.status != "approved" and profile.approved_data),
            "warnings": [],
            "effective_date": profile.effective_date if profile else None,
            "submitted_at": profile.submitted_at if profile else None,
        }
        # سابقه‌ی بیمه بر اساس نسخه‌ی تأییدشده (سابقه‌ی پیش از استخدامِ اعلام‌شده)، وگرنه نسخه‌ی جاری
        ins_data = (profile.approved_data or (self.snapshot(profile) if status != "none" else None)) if profile else None
        ins = insurance_info(settings, emp, profile, ins_data, as_of)
        item["insurance_days"], item["insurance_source"] = ins["days"], ins["source_label"]
        if profile is None or status == "none":
            return item
        current = self.snapshot(profile)
        item["marital_status"] = current["marital_status"]
        item["sons"] = sum(1 for m in current["members"] if m["member_type"] == "son")
        item["daughters"] = sum(1 for m in current["members"] if m["member_type"] == "daughter")
        item["warnings"] = rules.warnings(current, settings, today)
        if profile.approved_data:
            ev = rules.evaluate(profile.approved_data, emp.gender, ins["days"], settings, as_of)
            item["marriage_eligible"] = ev["marriage"]["eligible"]
            item["eligible_children"] = ev["child"]["eligible_count"]
        return item

    async def detail(self, profile_id: int, site_ids: set[int] | None, as_of: tuple[int, int, int]) -> dict | None:
        found = await self.get_profile(profile_id, site_ids)
        if found is None:
            return None
        profile, emp = found
        settings = await self.get_settings()
        current = self.snapshot(profile)
        logs = (
            await self.db.execute(
                select(FamilyChangeLog).where(FamilyChangeLog.profile_id == profile.id).order_by(FamilyChangeLog.id.desc())
            )
        ).scalars().all()
        site = await self.db.get(Site, emp.site_id)
        ins_current = insurance_info(settings, emp, profile, current, as_of)
        ins_approved = insurance_info(settings, emp, profile, profile.approved_data, as_of) if profile.approved_data else None
        return {
            "employee": {
                "id": emp.id,
                "personnel_code": emp.personnel_code,
                "first_name": emp.first_name,
                "last_name": emp.last_name,
                "gender": emp.gender,
                "national_code": emp.national_code,
                "hire_date_jalali": emp.hire_date_jalali,
                "site_name": site.name if site else None,
            },
            "profile": self.profile_out(profile, settings),
            "insurance_days": profile.insurance_days,
            "insurance": ins_current,  # سابقه‌ی بیمه با منبع و اجزا (نسخه‌ی جاری)
            "insurance_approved": ins_approved,
            "hr_note": profile.hr_note,
            "approved_data": profile.approved_data,
            "evaluation_current": rules.evaluate(current, emp.gender, ins_current["days"], settings, as_of),
            "evaluation_approved": rules.evaluate(profile.approved_data, emp.gender, ins_approved["days"], settings, as_of)
            if profile.approved_data
            else None,
            "missing_documents": rules.missing_documents(settings, current, today_jalali()) if profile.status != "draft" else [],
            "warnings": rules.warnings(current, settings, today_jalali()),
            "logs": [
                {
                    "id": log.id,
                    "action": log.action,
                    "action_label": ACTION_LABELS.get(log.action, log.action),
                    "actor_name": log.actor_name,
                    "note": log.note,
                    "effective_date": log.effective_date,
                    "created_at": log.created_at,
                }
                for log in logs
            ],
        }

    async def approve(self, profile_id: int, site_ids, user: User, effective_date: str | None, note: str | None):
        found = await self.get_profile(profile_id, site_ids)
        if found is None:
            return None
        profile, emp = found
        if profile.status not in ("pending", "returned", "rejected"):
            raise FamilyError("فقط پرونده‌ی ثبت‌شده‌ای که تأیید نشده قابل تأیید است.")
        if profile.marital_status is None:
            raise FamilyError("پرسنل هنوز فرم را ثبت نکرده است.")
        eff = rules.parse_jalali(effective_date) if effective_date else today_jalali()
        if eff is None:
            raise FamilyError("تاریخ اثر نامعتبر است.")
        snap = self.snapshot(profile)
        # مدارک بایگانی‌شده‌ی نسخه‌ی تأییدشده‌ی قبلی دیگر لازم نیستند
        archived = [d.id for d in profile.documents if d.archived]
        if archived:
            await self.db.execute(delete(FamilyDocument).where(FamilyDocument.id.in_(archived)))
        profile.status = "approved"
        profile.approved_data = snap
        profile.approved_at = datetime.now(timezone.utc)
        profile.effective_date = rules.format_jalali(eff)
        profile.reviewed_at = profile.approved_at
        profile.reviewed_by_id = user.id
        profile.review_note = (note or "").strip() or None
        self._log(profile, "approved", user, note=profile.review_note, effective_date=profile.effective_date, snapshot=snap)
        notice_id = await self._notify(
            emp.id,
            user,
            "مشخصات خانوادگی تأیید شد",
            f"مشخصات خانوادگی شما توسط واحد منابع انسانی تأیید شد (تاریخ اثر: {profile.effective_date})."
            + (f"\n{profile.review_note}" if profile.review_note else ""),
        )
        await self.db.commit()
        return notice_id

    async def reject(self, profile_id: int, site_ids, user: User, note: str, action: str):
        """action=rejected (رد) یا returned (بازگشت برای ویرایش، حتی وقتی پرونده قفل است)."""
        found = await self.get_profile(profile_id, site_ids)
        if found is None:
            return None
        profile, emp = found
        if profile.status == "draft":
            raise FamilyError("پرسنل هنوز فرم را ثبت نکرده است.")
        if action == "rejected" and profile.status != "pending":
            raise FamilyError("فقط پرونده‌ی «در انتظار بررسی» قابل رد است.")
        note = note.strip()
        profile.status = action
        profile.reviewed_at = datetime.now(timezone.utc)
        profile.reviewed_by_id = user.id
        profile.review_note = note
        self._log(profile, action, user, note=note)
        title = "مشخصات خانوادگی تأیید نشد" if action == "rejected" else "مشخصات خانوادگی نیاز به ویرایش دارد"
        notice_id = await self._notify(
            emp.id, user, title, f"{note}\nلطفاً از پنل کاربری، بخش «مشخصات خانوادگی»، فرم را اصلاح و دوباره ثبت کنید."
        )
        await self.db.commit()
        return notice_id

    async def update_hr_fields(self, employee_id: int, site_ids, user: User, payload: dict) -> dict | None:
        emp = await self.db.get(Employee, employee_id)
        if emp is None or (site_ids is not None and emp.site_id not in site_ids):
            return None
        profile = await self._get_or_create_shell(emp)
        changes = []
        if payload.get("clear_insurance_days"):
            profile.insurance_days = None
            changes.append("سابقه بیمه: پاک شد")
        elif payload.get("insurance_days") is not None:
            profile.insurance_days = int(payload["insurance_days"])
            changes.append(f"سابقه بیمه: {profile.insurance_days} روز")
        if payload.get("clear_hr_prior_insurance_days"):
            profile.hr_prior_insurance_days = None
            changes.append("سابقه‌ی قبلی (منابع انسانی): پاک شد")
        elif payload.get("hr_prior_insurance_days") is not None:
            profile.hr_prior_insurance_days = int(payload["hr_prior_insurance_days"])
            changes.append(f"سابقه‌ی قبلی (منابع انسانی): {profile.hr_prior_insurance_days} روز")
        if payload.get("hr_note") is not None:
            profile.hr_note = payload["hr_note"].strip() or None
            changes.append("یادداشت داخلی")
        if changes:
            self._log(profile, "hr_fields", user, note="، ".join(changes))
        await self.db.commit()
        return {
            "profile_id": profile.id,
            "insurance_days": profile.insurance_days,
            "hr_prior_insurance_days": profile.hr_prior_insurance_days,
            "hr_note": profile.hr_note,
        }

    async def import_insurance_days(
        self, site_ids: set[int] | None, user: User, content: bytes, mode: str, site_id: int | None = None
    ) -> dict:
        """
        ورود گروهی سابقه‌ی بیمه از Excel: ستون «کد پرسنلی» و ستون روز (عنوان شامل «روز» یا «سابقه»؛
        بدون عنوان مناسب: ستون اول و دوم). mode=prior → «سابقه‌ی پیش از استخدام» (با محاسبه‌ی خودکار جمع می‌شود)،
        mode=total → «کل سابقه» (جایگزین محاسبه). سلول خالی نادیده گرفته می‌شود.
        خروجی: {updated, not_found, invalid, ambiguous}
        """
        from openpyxl import load_workbook

        if mode not in ("prior", "total"):
            raise FamilyError("نوع ورود نامعتبر است.")
        def _read_rows() -> list:
            """خواندن همگام Excel با openpyxl (در Thread جدا تا حلقه‌ی async معطل نشود)."""
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            return list(wb.active.iter_rows(values_only=True))

        try:
            rows = await asyncio.to_thread(_read_rows)
        except Exception as e:  # noqa: BLE001
            raise FamilyError("فایل Excel قابل خواندن نیست (فقط xlsx).") from e
        if not rows:
            raise FamilyError("فایل خالی است.")
        header = [normalize_search_text(str(c or "")) for c in rows[0]]
        code_col = next((i for i, h in enumerate(header) if "کد" in h or "پرسنلی" in h), None)
        days_col = next((i for i, h in enumerate(header) if "روز" in h or "سابقه" in h), None)
        start = 1
        if code_col is None or days_col is None:
            code_col, days_col, start = 0, 1, 0  # بدون سطر عنوان

        q = select(Employee)
        if site_ids is not None:
            q = q.where(Employee.site_id.in_(site_ids))
        if site_id is not None:
            q = q.where(Employee.site_id == site_id)
        by_code: dict[str, list[Employee]] = {}
        for emp in (await self.db.execute(q)).scalars().all():
            by_code.setdefault(normalize_search_text(emp.personnel_code), []).append(emp)

        updated, not_found, invalid, ambiguous = 0, [], [], []
        label = "سابقه‌ی قبلی (منابع انسانی)" if mode == "prior" else "کل سابقه بیمه"
        for idx, row in enumerate(rows[start:], start=start + 1):
            if row is None or len(row) <= max(code_col, days_col):
                continue
            code = normalize_search_text(str(row[code_col] or "")).removesuffix(".0")
            raw_days = normalize_search_text(str(row[days_col] if row[days_col] is not None else "")).removesuffix(".0")
            if not code or not raw_days:
                continue
            if not raw_days.isdigit() or int(raw_days) > 20000:
                invalid.append(f"ردیف {idx}: {code}")
                continue
            emps = by_code.get(code) or []
            if not emps:
                not_found.append(code)
                continue
            if len(emps) > 1:
                ambiguous.append(code)  # کد تکراری در چند سایت: با انتخاب سایت دوباره وارد کنید
                continue
            profile = await self._get_or_create_shell(emps[0])
            if mode == "prior":
                profile.hr_prior_insurance_days = int(raw_days)
            else:
                profile.insurance_days = int(raw_days)
            self._log(profile, "hr_fields", user, note=f"ورود از Excel — {label}: {int(raw_days)} روز")
            updated += 1
        await self.db.commit()
        return {"updated": updated, "not_found": not_found, "invalid": invalid, "ambiguous": ambiguous}

    async def export_xlsx(self, site_ids: set[int] | None, as_of: tuple[int, int, int]) -> bytes:
        """
        دو برگه: «خلاصه» (یک سطر برای هر پرسنل با شمول بر اساس نسخه‌ی تأییدشده) و
        «اعضا» (همسر و فرزندان نسخه‌ی تأییدشده با علت شمول/عدم شمول هر فرزند).
        """
        from openpyxl import Workbook

        settings = await self.get_settings()
        data = await self.list_profiles(site_ids, None, None, None, 1, 100000, as_of)
        emp_ids = [i["employee_id"] for i in data["items"]]
        profiles: dict[int, FamilyProfile] = {}
        genders: dict[int, int | None] = {}
        if emp_ids:
            res = await self.db.execute(select(FamilyProfile).where(FamilyProfile.employee_id.in_(emp_ids)))
            profiles = {p.employee_id: p for p in res.scalars().all()}
            gres = await self.db.execute(select(Employee.id, Employee.gender).where(Employee.id.in_(emp_ids)))
            genders = dict(gres.all())

        yes_no = {True: "بله", False: "خیر", None: "—"}

        def _build() -> bytes:
            """ساخت Workbook — کاملاً همگام (openpyxl + rules.evaluate)؛ همه‌ی داده‌ها قبلاً با await واکشی شده‌اند."""
            wb = Workbook()
            ws = wb.active
            ws.title = "خلاصه"
            ws.sheet_view.rightToLeft = True
            ws.append([f"تاریخ مبنا: {rules.format_jalali(as_of)} — شمول بر اساس آخرین نسخه‌ی تأییدشده و تنظیمات فعلی"])
            ws.append([
                "کد پرسنلی", "نام", "نام خانوادگی", "سایت", "واحد", "وضعیت پرونده", "وضعیت تاهل", "سرپرست خانوار",
                "تعداد پسر", "تعداد دختر", "سابقه بیمه (روز)", "منبع سابقه بیمه", "مشمول حق تاهل", "علت", "فرزندان واجد شرایط حق اولاد",
                "تاریخ اثر", "تغییرات تأییدنشده", "هشدارها",
            ])  # fmt: skip
            members_rows = []
            for item in data["items"]:
                p = profiles.get(item["employee_id"])
                approved = p.approved_data if p else None
                # همان سابقه‌ی بیمه‌ی فهرست (خودکار از تاریخ استخدام + سابقه‌ی قبلی، یا اصلاح منابع انسانی)
                ev = (
                    rules.evaluate(approved, genders.get(item["employee_id"]), item["insurance_days"], settings, as_of)
                    if approved
                    else None
                )
                src = approved or {}
                ws.append([
                    item["personnel_code"], item["first_name"], item["last_name"], item["site_name"] or "",
                    item["department_name"] or "", rules.PROFILE_STATUSES.get(item["status"], "ثبت نشده"),
                    rules.MARITAL_STATUSES.get(src.get("marital_status"), "—"), yes_no[src.get("is_head_of_household")],
                    sum(1 for m in src.get("members", []) if m["member_type"] == "son"),
                    sum(1 for m in src.get("members", []) if m["member_type"] == "daughter"),
                    item["insurance_days"] if item["insurance_days"] is not None else "", item["insurance_source"] or "",
                    yes_no[ev["marriage"]["eligible"]] if ev else "—", ev["marriage"]["reason"] if ev else "",
                    ev["child"]["eligible_count"] if ev else "", item["effective_date"] or "",
                    "بله" if item["has_pending_changes"] else "", "؛ ".join(item["warnings"]),
                ])  # fmt: skip
                if approved:
                    kid_by_index = {k["member_index"]: k for k in ev["child"]["children"]} if ev else {}
                    for idx, m in enumerate(approved.get("members", [])):
                        k = kid_by_index.get(idx)
                        members_rows.append([
                            item["personnel_code"], f"{item['first_name']} {item['last_name']}",
                            rules.MEMBER_TYPES.get(m["member_type"], m["member_type"]), m["first_name"], m["last_name"],
                            m.get("national_id") or "", m.get("birth_date") or "", k["age"] if k and k["age"] is not None else "",
                            rules.RELATIONS.get(m.get("relation"), ""), yes_no[m.get("is_disabled")],
                            yes_no[m.get("is_student")] if m["member_type"] != "spouse" else "",
                            m.get("student_cert_expiry") or "",
                            yes_no[m.get("is_married")] if m["member_type"] == "daughter" else "",
                            yes_no[m.get("is_employed")],
                            (yes_no[k["eligible"]] if k else "") if m["member_type"] != "spouse" else "",
                            (k["reason"] if k else "") if m["member_type"] != "spouse" else "",
                        ])  # fmt: skip
            ws2 = wb.create_sheet("اعضا")
            ws2.sheet_view.rightToLeft = True
            ws2.append([
                "کد پرسنلی", "پرسنل", "نسبت", "نام", "نام خانوادگی", "کد ملی", "تاریخ تولد", "سن", "نوع فرزند",
                "از کار افتاده", "در حال تحصیل", "اعتبار گواهی تحصیل", "ازدواج کرده", "شاغل", "واجد شرایط حق اولاد", "علت",
            ])  # fmt: skip
            for r in members_rows:
                ws2.append(r)
            for sheet in (ws, ws2):
                for col in sheet.columns:
                    sheet.column_dimensions[col[0].column_letter].width = 16
            out = io.BytesIO()
            wb.save(out)
            return out.getvalue()

        # ساخت Excel در Thread جدا تا حلقه‌ی async معطل نشود
        return await asyncio.to_thread(_build)

    async def cleanup_unlinked_documents(self) -> int:
        """مدارکی که بیش از ۷۲ ساعت پیش آپلود شده و هرگز در فرم ثبت نشده‌اند پاک می‌شوند."""
        cutoff = datetime.now(timezone.utc) - timedelta(hours=UNLINKED_MAX_AGE_HOURS)
        result = await self.db.execute(
            delete(FamilyDocument).where(FamilyDocument.linked.is_(False), FamilyDocument.uploaded_at < cutoff)
        )
        await self.db.commit()
        return result.rowcount or 0


def _approved_doc_ids(approved: dict | None) -> set[int]:
    """شناسه‌ی همه‌ی مدارکی که در نسخه‌ی تأییدشده ارجاع شده‌اند."""
    if not approved:
        return set()
    ids = {d.get("id") for d in approved.get("docs") or []}
    for m in approved.get("members") or []:
        ids |= {d.get("id") for d in m.get("docs") or []}
    return {i for i in ids if isinstance(i, int)}
