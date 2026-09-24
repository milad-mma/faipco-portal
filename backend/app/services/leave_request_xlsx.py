"""
تولید فایل Excel از گزارش مدیریتی درخواست‌های مرخصی/ماموریت (endpoint export).

ورودی همان لیست دیکشنری‌های نرمالایزشده‌ی list_all_for_site است؛ خروجی بایت‌های xlsx
با قالب گزارش‌های ارزیابی عملکرد (RTL، سرستون آبی) و تاریخ‌های شمسی.
"""
from __future__ import annotations

import io

import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# استایل سرستون‌ها: پس‌زمینه آبی تیره، متن سفید بولد، راست‌چین
_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_RTL_ALIGNMENT = Alignment(horizontal="right", readingOrder=2)

# برچسب فارسی وضعیت‌های نرمالایزشده
_STATUS_LABELS = {"pending": "در حال بررسی", "approved": "تائید شده", "rejected": "رد شده", "cancelled": "ابطال شده"}

# ستون‌های خروجی به ترتیب: (عنوان سرستون، عرض ستون)
_COLUMNS = [
    ("نام و نام خانوادگی", 26),
    ("کد پرسنلی", 14),
    ("واحد", 22),
    ("نوع درخواست", 24),
    ("توضیحات", 34),
    ("تاریخ ثبت", 14),
    ("تاریخ شروع", 14),
    ("تاریخ پایان", 14),
    ("ساعت شروع", 12),
    ("ساعت پایان", 12),
    ("مدت", 16),
    ("وضعیت", 14),
    ("نظر تأییدکننده", 34),
]


def _to_jalali(value) -> str:
    """یک date/datetime میلادی را به رشته شمسی YYYY/MM/DD تبدیل می‌کند؛ مقدار خالی یا نامعتبر، رشته خالی."""
    if not value:
        return ""
    try:
        d = value.date() if hasattr(value, "date") else value  # datetime به date
        j = jdatetime.date.fromgregorian(date=d)
        return f"{j.year}/{j.month:02d}/{j.day:02d}"
    except Exception:
        return ""


def _format_compact_time(compact) -> str:
    """ساعت فشرده HHMM (مثلاً 1236) را به «12:36» تبدیل می‌کند؛ None به رشته خالی."""
    if compact is None:
        return ""
    return f"{compact // 100:02d}:{compact % 100:02d}"


def _format_duration(item: dict) -> str:
    """مدت یک درخواست را مثل فرانت‌اند می‌سازد: ساعتی -> ساعت و دقیقه؛ روزانه -> تعداد روز."""
    start_hour, end_hour = item.get("start_hour"), item.get("end_hour")
    # درخواست ساعتی: اختلاف دقیقه‌ای ساعت شروع و پایان
    if start_hour is not None and end_hour is not None:
        diff = (end_hour // 100 * 60 + end_hour % 100) - (start_hour // 100 * 60 + start_hour % 100)
        if diff <= 0:
            return ""
        hours, minutes = divmod(diff, 60)
        if hours and minutes:
            return f"{hours} ساعت و {minutes} دقیقه"
        return f"{hours} ساعت" if hours else f"{minutes} دقیقه"
    # درخواست روزانه: تعداد روزهای بازه (هر دو سر بازه شامل)
    start_date, end_date = item.get("start_date"), item.get("end_date")
    if start_date and end_date:
        days = (end_date - start_date).days + 1
        return f"{days} روز" if days > 0 else ""
    return ""


def build_leave_requests_xlsx(items: list[dict], site_name: str = "") -> bytes:
    """
    از لیست درخواست‌های نرمالایزشده یک فایل xlsx می‌سازد و بایت‌های آن را برمی‌گرداند.
    ورودی: آیتم‌های خروجی list_all_for_site و نام سایت (اختیاری، در سطر اول نوشته می‌شود).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "درخواست‌ها"
    ws.sheet_view.rightToLeft = True

    # سطر عنوان سایت و یک سطر خالی قبل از جدول
    if site_name:
        ws.append(["سایت", site_name])
        ws.append([])

    # سرستون‌ها با استایل
    header_row_index = ws.max_row + 1
    ws.append([title for title, _ in _COLUMNS])
    for col in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(row=header_row_index, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _RTL_ALIGNMENT

    # یک سطر به‌ازای هر درخواست، به همان ترتیب _COLUMNS
    for item in items:
        ws.append(
            [
                item.get("requester_name") or "",
                item.get("emp_no") or "",
                item.get("requester_department") or "",
                item.get("type_title") or "",
                item.get("description") or "",
                _to_jalali(item.get("submitted_at")),
                _to_jalali(item.get("start_date")),
                _to_jalali(item.get("end_date")),
                _format_compact_time(item.get("start_hour")),
                "" if item.get("is_forgotten_punch") else _format_compact_time(item.get("end_hour")),  # تردد فراموش‌شده ساعت پایان ندارد
                "" if item.get("is_forgotten_punch") else _format_duration(item),
                "در انتظار منابع انسانی"  # وضعیت ویژه تردد فراموش‌شده بعد از تأیید سرپرست
                if item.get("awaiting_hr")
                else _STATUS_LABELS.get(item.get("status"), item.get("status") or ""),
                item.get("manager_idea") or "",
            ]
        )

    # عرض ستون‌ها
    for i, (_, width) in enumerate(_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # خروجی در حافظه، بدون فایل موقت
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
